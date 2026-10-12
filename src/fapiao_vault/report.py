"""报表输出：JSON 台账、CSV、Excel 与自包含 HTML 报告。

HTML 报告把样式和数据全部内联在单个文件里，不引用任何 CDN，
双击即可离线打开——既符合"零上传"的定位，也方便直接发给财务同事。
"""

from __future__ import annotations

import csv
import html
import json
from collections.abc import Iterable
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from .models import Invoice

# ------------------------------------------------------------------ 汇总


def _sum(values: Iterable[Decimal | None]) -> Decimal:
    total = Decimal("0.00")
    for v in values:
        if v is not None:
            total += v
    return total.quantize(Decimal("0.01"))


def build_summary(result) -> dict[str, Any]:
    """把 ScanResult 汇总成给报表用的字典。"""
    invoices: list[Invoice] = result.invoices
    ok = result.ok_invoices

    by_month: dict[str, dict[str, Any]] = {}
    by_seller: dict[str, dict[str, Any]] = {}
    by_type: dict[str, dict[str, Any]] = {}

    for inv in ok:
        for bucket, key in (
            (by_month, inv.month),
            (by_seller, inv.seller_name or "未知销售方"),
            (by_type, inv.invoice_type or "未知票种"),
        ):
            entry = bucket.setdefault(key, {"count": 0, "total": Decimal("0.00")})
            entry["count"] += 1
            entry["total"] += inv.total or Decimal("0.00")

    def dump(bucket: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
        rows = [
            {"key": k, "count": v["count"], "total": str(v["total"].quantize(Decimal("0.01")))}
            for k, v in bucket.items()
        ]
        rows.sort(key=lambda r: Decimal(r["total"]), reverse=True)
        return rows

    issues = [
        {
            "file": Path(inv.source_file).name,
            "confidence": inv.confidence,
            "problems": inv.issues,
        }
        for inv in invoices
        if inv.issues
    ]

    duplicates = [
        {
            "uid": g.uid,
            "severity": g.severity,
            "reason": g.reason,
            "files": [Path(i.source_file).name for i in g.invoices],
            "total": str(g.invoices[0].total) if g.invoices[0].total is not None else "",
            "seller": g.invoices[0].seller_name,
            "date": g.invoices[0].invoice_date,
        }
        for g in result.duplicates
    ]

    cross_period = [
        {
            "uid": inv.uid,
            "invoice_number": inv.invoice_number,
            "month": inv.month,
            "history_month": history_month,
            "total": str(inv.total) if inv.total is not None else "",
            "seller": inv.seller_name,
            "file": Path(inv.source_file).name,
        }
        for inv, history_month in getattr(result, "cross_period", [])
    ]

    return {
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "totals": {
            "scanned": result.scanned_count,
            "ok": len(ok),
            "failed": len(result.failed_invoices),
            "duplicate_groups": len(result.duplicates),
            "high_risk": len([d for d in result.duplicates if d.severity == "high"]),
            "cross_period": len(cross_period),
            "total": str(_sum(i.total for i in ok)),
            "tax": str(_sum(i.tax for i in ok)),
            "amount": str(_sum(i.amount for i in ok)),
        },
        "by_month": dump(by_month),
        "by_seller": dump(by_seller)[:20],
        "by_type": dump(by_type),
        "duplicates": duplicates,
        "cross_period": cross_period,
        "issues": issues,
        "rows": [inv.to_dict() for inv in invoices],
    }


# ------------------------------------------------------------------ JSON


def write_json(summary: dict[str, Any], path: Path | str) -> Path:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return p


# ------------------------------------------------------------------ CSV

CSV_HEADERS = [
    ("invoice_number", "发票号码"),
    ("invoice_code", "发票代码"),
    ("invoice_date", "开票日期"),
    ("invoice_type", "票种"),
    ("seller_name", "销售方"),
    ("seller_tax_id", "销售方税号"),
    ("buyer_name", "购买方"),
    ("buyer_tax_id", "购买方税号"),
    ("amount", "金额"),
    ("tax", "税额"),
    ("total", "价税合计"),
    ("check_code", "校验码"),
    ("confidence", "置信度"),
    ("source_file", "来源文件"),
    ("issues", "问题"),
]


def write_csv(summary: dict[str, Any], path: Path | str) -> Path:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    # utf-8-sig 让 Excel 双击打开不乱码
    with p.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow([label for _, label in CSV_HEADERS])
        for row in summary["rows"]:
            writer.writerow(
                [
                    "；".join(row.get("issues", [])) if key == "issues" else row.get(key, "")
                    for key, _ in CSV_HEADERS
                ]
            )
    return p


# ------------------------------------------------------------------ Excel


def write_xlsx(summary: dict[str, Any], path: Path | str) -> Path | None:
    """导出 Excel。未安装 openpyxl 时返回 None（调用方降级为 CSV）。"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None

    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F6FEB")
    warn_fill = PatternFill("solid", fgColor="FFF3CD")
    high_fill = PatternFill("solid", fgColor="FFD7D5")

    # --- Sheet 1：发票明细 ---
    ws = wb.active
    ws.title = "发票明细"
    ws.append([label for _, label in CSV_HEADERS])
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row in summary["rows"]:
        ws.append(
            [
                "；".join(row.get("issues", [])) if key == "issues" else row.get(key, "")
                for key, _ in CSV_HEADERS
            ]
        )
    widths = [24, 14, 12, 22, 30, 22, 30, 22, 12, 12, 12, 24, 10, 46, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    # 有问题的行标黄
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=15).value:
            for c in range(1, len(CSV_HEADERS) + 1):
                ws.cell(row=r, column=c).fill = warn_fill

    # --- Sheet 2：报销汇总 ---
    ws2 = wb.create_sheet("报销汇总")
    ws2.append(["统计维度", "项目", "张数", "金额（价税合计）"])
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
    for item in summary["by_month"]:
        ws2.append(["按月份", item["key"], item["count"], item["total"]])
    for item in summary["by_seller"]:
        ws2.append(["按销售方", item["key"], item["count"], item["total"]])
    total_row = [
        "总计",
        "",
        summary["totals"]["ok"],
        summary["totals"]["total"],
    ]
    ws2.append(total_row)
    for cell in ws2[ws2.max_row]:
        cell.font = Font(bold=True)
    for i, w in enumerate([14, 40, 10, 18], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # --- Sheet 3：重复发票 ---
    ws3 = wb.create_sheet("重复发票")
    ws3.append(["风险等级", "标识", "原因", "涉及文件", "金额", "销售方", "开票日期"])
    for cell in ws3[1]:
        cell.font = header_font
        cell.fill = header_fill
    for d in summary["duplicates"]:
        ws3.append(
            [
                "高风险" if d["severity"] == "high" else "提示",
                d["uid"],
                d["reason"],
                " | ".join(d["files"]),
                d["total"],
                d["seller"],
                d["date"],
            ]
        )
        if d["severity"] == "high":
            for c in range(1, 8):
                ws3.cell(row=ws3.max_row, column=c).fill = high_fill

    # 与历史台账重号（跨期重复报销）
    for cp in summary.get("cross_period", []):
        ws3.append(
            [
                "高风险(跨期)",
                cp["invoice_number"] or cp["uid"],
                f"台账月份 {cp['history_month']}，本次 {cp['month']}，疑似跨期重复报销",
                cp["file"],
                cp["total"],
                cp["seller"],
                "",
            ]
        )
        for c in range(1, 8):
            ws3.cell(row=ws3.max_row, column=c).fill = high_fill
    for i, w in enumerate([10, 26, 46, 60, 12, 28, 12], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    wb.save(p)
    return p


# ------------------------------------------------------------------ HTML

HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>发票整理报告 · FapiaoVault</title>
<style>
  :root{
    --bg:#f6f7f9; --card:#fff; --fg:#1a1d21; --muted:#6b7280;
    --border:#e5e7eb; --accent:#1f6feb; --ok:#0f9d58; --warn:#b7791f; --danger:#d93025;
    --radius:12px;
  }
  @media (prefers-color-scheme: dark){
    :root{ --bg:#0f1115; --card:#171a21; --fg:#e6e8eb; --muted:#9aa1ac;
           --border:#262b35; --accent:#4c8dff; }
  }
  *{box-sizing:border-box}
  body{margin:0;padding:28px 20px 60px;background:var(--bg);color:var(--fg);
       font:15px/1.6 -apple-system,"Segoe UI","Microsoft YaHei",system-ui,sans-serif}
  .wrap{max-width:1180px;margin:0 auto}
  h1{font-size:24px;margin:0 0 4px}
  .sub{color:var(--muted);font-size:13px;margin-bottom:22px}
  .cards{display:grid;grid-template-columns:repeat(6,1fr);gap:14px;margin-bottom:22px}
  .card{background:var(--card);border:1px solid var(--border);border-radius:var(--radius);padding:16px 18px}
  .card.primary{grid-column:span 2}
  .card .k{font-size:12px;color:var(--muted);margin-bottom:6px}
  .card .v{font-size:26px;font-weight:600;font-variant-numeric:tabular-nums}
  .card .v.sm{font-size:19px}
  .card .v.ok{color:var(--ok)}
  .card .v.warn{color:var(--warn)}
  .card .v.danger{color:var(--danger)}
  @media (max-width: 920px){
    .cards{grid-template-columns:repeat(2,1fr)}
    .card.primary{grid-column:span 2}
  }
  .section{background:var(--card);border:1px solid var(--border);border-radius:var(--radius);
           margin-bottom:18px;overflow:hidden}
  .section > h2{font-size:15px;margin:0;padding:13px 18px;border-bottom:1px solid var(--border)}
  .section.alert > h2{color:var(--danger);box-shadow:inset 3px 0 0 var(--danger)}
  .section .body{padding:10px 18px}
  table{width:100%;border-collapse:collapse;font-size:13.5px}
  th,td{text-align:left;padding:9px 10px;border-bottom:1px solid var(--border);white-space:nowrap}
  th{color:var(--muted);font-weight:600;font-size:12px;background:transparent;position:sticky;top:0;background:var(--card)}
  td.wrapcell{white-space:normal;max-width:420px}
  tr:last-child td{border-bottom:none}
  .num{text-align:right;font-variant-numeric:tabular-nums}
  .pill{display:inline-block;padding:1px 8px;border-radius:999px;font-size:11.5px;border:1px solid var(--border)}
  .pill.high{color:var(--danger);border-color:var(--danger)}
  .pill.medium{color:var(--warn);border-color:var(--warn)}
  .pill.ok{color:var(--ok);border-color:var(--ok)}
  .toolbar{display:flex;gap:10px;padding:12px 18px;border-bottom:1px solid var(--border);flex-wrap:wrap}
  input,select{background:var(--bg);color:var(--fg);border:1px solid var(--border);
                border-radius:8px;padding:7px 10px;font:inherit;font-size:13px}
  input{flex:1;min-width:180px}
  .empty{color:var(--muted);font-size:13px;padding:6px 0}
  .notice{border-left:3px solid var(--accent);padding:8px 12px;background:rgba(31,111,235,.06);
          border-radius:6px;font-size:13px;color:var(--muted);margin-bottom:18px}
  .scroll{max-height:560px;overflow:auto}
</style>
</head>
<body>
<div class="wrap">
  <h1>发票整理报告</h1>
  <div class="sub" id="subtitle"></div>
  <div class="notice">本报告由 FapiaoVault 在本地生成，未联网、未上传任何数据。所有金额来自发票文本层解析结果，提交前请与原始票据核对。</div>
  <div class="cards" id="cards"></div>
  <div id="sections"></div>
</div>
<script>
const DATA = __DATA_JSON__;

const fmt = (v) => v === null || v === undefined || v === "" ? "" : v;
const money = (v) => (v === "" || v == null) ? "—" : "¥" + Number(v).toLocaleString("zh-CN",{minimumFractionDigits:2,maximumFractionDigits:2});
const esc = (s) => String(s ?? "").replace(/[&<>"']/g, c => ({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#39;"}[c]));

function card(k, v, cls){
  const cardCls = cls && cls.includes("primary") ? "card primary" : "card";
  const vCls = (cls || "").split(" ").filter(c => c && c !== "primary").join(" ");
  return `<div class="${cardCls}"><div class="k">${esc(k)}</div><div class="v ${vCls}">${v}</div></div>`;
}

function render(){
  const t = DATA.totals;
  document.getElementById("subtitle").textContent =
    `扫描时间：${DATA.generated_at}　·　识别成功 ${t.ok} 张 / 共扫描 ${t.scanned} 个文件`;

  document.getElementById("cards").innerHTML = [
    card("价税合计总额", money(t.total), "primary"),
    card("其中税额", money(t.tax), "sm"),
    card("识别成功", `${t.ok} 张`, "sm ok"),
    card("重复发票", `${t.duplicate_groups} 组` + (t.high_risk ? `（高风险 ${t.high_risk}）` : ""),
         t.duplicate_groups > 0 ? "sm danger" : "sm ok"),
    card("需人工处理", `${t.failed} 个文件`, t.failed > 0 ? "sm warn" : "sm ok"),
  ].join("");

  const sec = [];
  sec.push(crossSection());
  sec.push(dupSection());
  sec.push(tableSection());
  sec.push(groupSection("按月份汇总", DATA.by_month, "月份"));
  sec.push(groupSection("按销售方汇总（前 20）", DATA.by_seller, "销售方"));
  sec.push(groupSection("按票种汇总", DATA.by_type, "票种"));
  sec.push(issueSection());
  document.getElementById("sections").innerHTML = sec.join("");
  bindTable();
}

function crossSection(){
  if(!DATA.cross_period || !DATA.cross_period.length) return "";
  const rows = DATA.cross_period.map(c => `<tr>
    <td><span class="pill high">高风险</span></td>
    <td>${esc(c.invoice_number || c.uid)}</td>
    <td class="wrapcell">台账月份 ${esc(c.history_month)}，本次 ${esc(c.month)} —— 疑似跨期重复报销</td>
    <td class="num">${money(c.total)}</td>
    <td class="wrapcell">${esc(c.file)}</td>
  </tr>`).join("");
  return `<div class="section alert"><h2>与历史台账重号（${DATA.cross_period.length}，疑似跨期重复报销）</h2>
    <div class="body"><table><thead><tr>
      <th>等级</th><th>发票标识</th><th>说明</th><th class="num">金额</th><th>文件</th>
    </tr></thead><tbody>${rows}</tbody></table></div></div>`;
}

function dupSection(){
  if(!DATA.duplicates.length) return "";
  const rows = DATA.duplicates.map(d => `<tr>
    <td><span class="pill ${d.severity === "high" ? "high" : "medium"}">${d.severity === "high" ? "高风险" : "提示"}</span></td>
    <td>${esc(d.uid)}</td>
    <td class="wrapcell">${esc(d.reason)}</td>
    <td class="num">${money(d.total)}</td>
    <td class="wrapcell">${d.files.map(esc).join("<br>")}</td>
  </tr>`).join("");
  return `<div class="section alert"><h2>重复发票警告（${DATA.duplicates.length} 组）</h2>
    <div class="body"><table><thead><tr>
      <th>等级</th><th>发票标识</th><th>原因</th><th class="num">金额</th><th>涉及文件</th>
    </tr></thead><tbody>${rows}</tbody></table></div></div>`;
}

function tableSection(){
  const rows = DATA.rows;
  return `<div class="section"><h2>发票明细（${rows.length}）</h2>
    <div class="toolbar">
      <input id="q" placeholder="搜索：发票号 / 销售方 / 文件名 / 日期">
      <select id="conf">
        <option value="">全部置信度</option>
        <option value="high">仅高置信度</option>
        <option value="medium">仅中等</option>
        <option value="low">仅有问题</option>
      </select>
    </div>
    <div class="scroll"><table id="detail"><thead><tr>
      <th>开票日期</th><th>发票号码</th><th>销售方</th><th class="num">金额</th>
      <th class="num">税额</th><th class="num">价税合计</th><th>票种</th><th>置信度</th><th>问题</th><th>来源文件</th>
    </tr></thead><tbody>
    ${rows.map(r => `<tr data-conf="${esc(r.confidence)}" data-search="${esc([r.invoice_number,r.seller_name,r.source_file,r.invoice_date,r.buyer_name].join(" ").toLowerCase())}">
      <td>${esc(fmt(r.invoice_date))}</td>
      <td>${esc(fmt(r.invoice_number))}</td>
      <td class="wrapcell">${esc(fmt(r.seller_name))}</td>
      <td class="num">${money(r.amount)}</td>
      <td class="num">${money(r.tax)}</td>
      <td class="num">${money(r.total)}</td>
      <td>${esc(fmt(r.invoice_type))}</td>
      <td><span class="pill ${r.confidence === "high" ? "ok" : (r.confidence === "low" ? "high" : "medium")}">${esc(r.confidence)}</span></td>
      <td class="wrapcell">${(r.issues||[]).map(esc).join("<br>")}</td>
      <td class="wrapcell">${esc(fmt(r.source_file))}</td>
    </tr>`).join("")}
    </tbody></table></div></div>`;
}

function groupSection(title, data, label){
  if(!data || !data.length) return `<div class="section"><h2>${esc(title)}</h2><div class="body"><div class="empty">暂无数据</div></div></div>`;
  return `<div class="section"><h2>${esc(title)}</h2><div class="body"><table><thead><tr>
    <th>${esc(label)}</th><th class="num">张数</th><th class="num">价税合计</th>
  </tr></thead><tbody>${data.map(d => `<tr><td class="wrapcell">${esc(d.key)}</td><td class="num">${d.count}</td><td class="num">${money(d.total)}</td></tr>`).join("")}</tbody></table></div></div>`;
}

function issueSection(){
  if(!DATA.issues.length) return "";
  return `<div class="section"><h2>需人工核对的问题（${DATA.issues.length}）</h2><div class="body"><table><thead><tr>
    <th>文件</th><th>置信度</th><th>问题</th>
  </tr></thead><tbody>${DATA.issues.map(i => `<tr>
    <td class="wrapcell">${esc(i.file)}</td>
    <td><span class="pill ${i.confidence === "low" ? "high" : "medium"}">${esc(i.confidence)}</span></td>
    <td class="wrapcell">${i.problems.map(esc).join("<br>")}</td>
  </tr>`).join("")}</tbody></table></div></div>`;
}

function bindTable(){
  const q = document.getElementById("q");
  const conf = document.getElementById("conf");
  if(!q) return;
  const rows = Array.from(document.querySelectorAll("#detail tbody tr"));
  function apply(){
    const kw = q.value.trim().toLowerCase();
    const c = conf.value;
    rows.forEach(tr => {
      const hitKw = !kw || tr.dataset.search.includes(kw);
      const hitC = !c || tr.dataset.conf === c;
      tr.style.display = (hitKw && hitC) ? "" : "none";
    });
  }
  q.addEventListener("input", apply);
  conf.addEventListener("change", apply);
}

render();
</script>
</body>
</html>
"""


def write_html(summary: dict[str, Any], path: Path | str, title: str = "发票整理报告") -> Path:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(summary, ensure_ascii=False).replace("</", "<\\/")
    doc = HTML_TEMPLATE.replace("__DATA_JSON__", payload)
    if title != "发票整理报告":
        doc = doc.replace("发票整理报告", html.escape(title))
    p.write_text(doc, encoding="utf-8")
    return p
